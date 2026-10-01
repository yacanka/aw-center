from io import BytesIO
from django.test import SimpleTestCase
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfWriter

from comparison.contracts import CompareError, ComparisonResult, Difference, resolve_options
from comparison.readers import read_document
from comparison.reports import write_report


class DocumentComparisonTests(SimpleTestCase):
    def test_word_preserves_text_and_table_content(self):
        document = Document()
        document.add_paragraph('Pressure  100')
        document.add_table(rows=1, cols=1).cell(0, 0).text = 'Table cell'
        content = BytesIO()
        document.save(content)
        blocks, warnings = read_document(content.getvalue(), 'word')
        self.assertIn('Pressure  100', [block.text for block in blocks])
        self.assertIn('Table cell', [block.text for block in blocks])
        self.assertEqual(warnings, [])

    def test_textless_pdf_is_not_reported_equal(self):
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        content = BytesIO()
        writer.write(content)
        with self.assertRaises(CompareError) as raised:
            read_document(content.getvalue(), 'pdf')
        self.assertEqual(raised.exception.code, 'COMPARE_TEXT_UNAVAILABLE')

    def test_both_reports_include_changes_and_neutralize_formula_text(self):
        import tempfile
        from pathlib import Path
        entries = [Difference('replace', '=1+1', '@SUM(A1)', 'Old P1', 'New P1', 'possible', .7)]
        for output in ('word', 'excel'):
            with self.subTest(output=output), tempfile.TemporaryDirectory() as directory:
                result = ComparisonResult('word', resolve_options({'output_type': output}), entries, 'text')
                path = Path(directory) / ('report.docx' if output == 'word' else 'report.xlsx')
                write_report(result, path)
                if output == 'excel':
                    workbook = load_workbook(path)
                    sheet = workbook['Differences']
                    text = [cell for row in sheet.iter_rows() for cell in row if cell.value]
                    self.assertTrue(any('@SUM(A1)' in str(cell.value) for cell in text))
                    self.assertTrue(all(cell.data_type != 'f' for cell in text))
                    self.assertEqual(workbook['Possible matches'].max_row, 2)
                else:
                    doc = Document(path)
                    text = '\n'.join(p.text for p in doc.paragraphs)
                    self.assertIn('=1+1', text)
                    self.assertIn('@SUM(A1)', text)
                    self.assertIn('Old P1', text)
