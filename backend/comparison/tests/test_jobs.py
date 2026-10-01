from datetime import timedelta
from io import BytesIO
import json
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from docx import Document
from openpyxl import load_workbook
from rest_framework.test import APIClient

from awcenter.job_executors import resolve_job_executor
from jobs.models import Job
from jobs.worker import claim_next_job, execute_claimed_job
from jobs.tests.base import JobTestCase
from comparison.tests.test_excel import book

BASE = '/api/tools/compare/'


def word(text):
    buffer = BytesIO()
    document = Document()
    document.add_paragraph(text)
    document.save(buffer)
    return buffer.getvalue()


def upload(content, name):
    return SimpleUploadedFile(name, content, content_type='application/octet-stream')


class ComparisonJobTests(JobTestCase):
    def enqueue(self, path, first, second, name='input.docx', params=None, key='comparison-test-key'):
        return self.client.post(BASE + path, {'first': upload(first, name), 'second': upload(second, name),
            'parameters': json.dumps(params or {})}, format='multipart', HTTP_IDEMPOTENCY_KEY=key)

    def run_job(self, response):
        self.assertEqual(response.status_code, 201, getattr(response, 'data', ''))
        job = Job.objects.get(pk=response.data['id'])
        claimed = claim_next_job('comparison-test-worker', [job.kind])
        execute_claimed_job(claimed, resolve_job_executor)
        job.refresh_from_db()
        self.assertEqual(job.status, 'succeeded', job.message)
        return job

    def test_document_job_and_idempotent_replay_download(self):
        old, new = word('Pressure 100'), word('Pressure 101')
        response = self.enqueue('jobs/', old, new, params={'output_type': 'excel'})
        replay = self.enqueue('jobs/', old, new, params={'output_type': 'excel'})
        self.assertEqual(replay.status_code, 200)
        self.assertEqual(response.data['id'], replay.data['id'])
        conflict = self.enqueue('jobs/', old, word('Other'), params={'output_type': 'excel'})
        self.assertEqual(conflict.status_code, 409)
        job = self.run_job(response)
        self.assertEqual(job.result_summary['replace'], 1)
        self.assertNotIn('Pressure', json.dumps(job.result_summary))
        result = self.client.get(f'/api/jobs/{job.id}/download/')
        self.assertEqual(result.status_code, 200)
        content = b''.join(result.streaming_content)
        sheet = load_workbook(BytesIO(content))['Differences']
        self.assertIn('Pressure 101', [cell.value for row in sheet for cell in row])

    def test_excel_inspection_then_report_uses_owned_snapshot(self):
        old = book([['ID', 'Name'], [1, 'Before']])
        new = book([['ID', 'Name'], [1, 'After']])
        inspection_job = self.run_job(self.enqueue('inspections/', old, new, 'input.xlsx'))
        inspection = self.client.get(BASE + f'inspections/{inspection_job.id}/')
        self.assertEqual(inspection.status_code, 200)
        self.assertFalse(inspection.data['matching']['requires_input'])
        response = self.client.post(BASE + 'jobs/', {'inspection_id': str(inspection_job.id),
            'parameters': {'output_type': 'excel'}}, format='json', HTTP_IDEMPOTENCY_KEY='compare-from-inspection')
        self.run_job(response)
        self.client.force_authenticate(self.other_user)
        self.assertEqual(self.client.get(BASE + f'inspections/{inspection_job.id}/').status_code, 404)
        self.assertEqual(self.client.post(BASE + 'jobs/', {'inspection_id': str(inspection_job.id)},
            format='json', HTTP_IDEMPOTENCY_KEY='other-owner-request').status_code, 404)

    def test_uncertain_excel_requires_new_inspection_with_choice(self):
        old = book([['Name'], ['A'], ['A']])
        new = book([['Name'], ['B'], ['B']])
        source = self.run_job(self.enqueue('inspections/', old, new, 'input.xlsx'))
        response = self.client.post(BASE + 'jobs/', {'inspection_id': str(source.id)}, format='json',
                                    HTTP_IDEMPOTENCY_KEY='ambiguous-compare')
        self.assertEqual(response.status_code, 400)
        revised = self.client.post(BASE + 'inspections/', {'inspection_id': str(source.id),
            'parameters': {'selection': {'matching': {'mode': 'position'}}}}, format='json',
            HTTP_IDEMPOTENCY_KEY='revise-inspection')
        job = self.run_job(revised)
        detail = self.client.get(BASE + f'inspections/{job.id}/')
        self.assertFalse(detail.data['matching']['requires_input'])

    def test_integrity_and_expired_inspections_rejected(self):
        content = book([['ID'], [1]])
        job = self.run_job(self.enqueue('inspections/', content, content, 'input.xlsx'))
        with job.output_file.open('wb') as handle:
            handle.write(b'{}')
        self.assertEqual(self.client.get(BASE + f'inspections/{job.id}/').status_code, 409)
        Job.objects.filter(pk=job.pk).update(completed_at=timezone.now() - timedelta(days=1000))
        self.assertEqual(self.client.get(BASE + f'inspections/{job.id}/').status_code, 404)

    def test_auth_csrf_mixed_formats_and_invalid_parameters(self):
        client = APIClient(enforce_csrf_checks=True)
        self.assertEqual(client.get(BASE + 'presets/').status_code, 403)
        client.force_login(self.user)
        self.assertEqual(client.post(BASE + 'jobs/', {}).status_code, 403)
        response = self.client.post(BASE + 'jobs/', {'first': upload(word('A'), 'input.docx'),
            'second': upload(book([['ID'], [1]]), 'input.xlsx')}, format='multipart',
            HTTP_IDEMPOTENCY_KEY='mixed-format-request')
        self.assertEqual(response.status_code, 400)
        self.assertEqual(Job.objects.count(), 0)

    def test_parser_failure_is_sanitized(self):
        response = self.enqueue('jobs/', word('A'), word('B'))
        job = Job.objects.get(pk=response.data['id'])
        with patch('comparison.executor.read_document', side_effect=RuntimeError('private-data')):
            execute_claimed_job(claim_next_job('comparison-error-worker'), resolve_job_executor)
        job.refresh_from_db()
        self.assertEqual(job.status, 'failed')
        self.assertNotIn('private-data', job.message)

    def test_all_document_and_report_pairs_execute(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        def pdf(text):
            writer = PdfWriter()
            page = writer.add_blank_page(width=300, height=300)
            font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                                     NameObject('/Subtype'): NameObject('/Type1'),
                                     NameObject('/BaseFont'): NameObject('/Helvetica')})
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
            stream = DecodedStreamObject()
            stream.set_data(f'BT /F1 12 Tf 20 240 Td ({text}) Tj ET'.encode())
            page[NameObject('/Contents')] = stream
            buffer = BytesIO()
            writer.write(buffer)
            return buffer.getvalue()
        for extension, create in [('docx', word), ('docm', word), ('pdf', pdf)]:
            for output in ('word', 'excel'):
                with self.subTest(extension=extension, output=output):
                    response = self.enqueue('jobs/', create('Pressure 100'), create('Pressure 101'), f'input.{extension}',
                                            {'output_type': output}, f'pair-{extension}-{output}')
                    job = self.run_job(response)
                    self.assertEqual(job.result_summary['replace'], 1)
                    self.assertTrue(job.output_name.endswith('.docx' if output == 'word' else '.xlsx'))

    def test_cancelled_job_is_not_claimed_and_corrupt_input_fails(self):
        response = self.enqueue('jobs/', word('A'), word('B'))
        job = Job.objects.get(pk=response.data['id'])
        self.client.post(f'/api/jobs/{job.id}/cancel/')
        self.assertIsNone(claim_next_job('cancel-test', ['comparison.compare']))
        response = self.enqueue('jobs/', word('A'), word('B'), key='corrupt-input-job')
        job = Job.objects.get(pk=response.data['id'])
        with job.input_file.open('wb') as handle:
            handle.write(b'corrupt')
        execute_claimed_job(claim_next_job('integrity-test'), resolve_job_executor)
        job.refresh_from_db()
        self.assertEqual(job.error_code, 'JOB_INPUT_CORRUPT')

    def test_unknown_nested_selection_fields_are_not_persisted(self):
        content = book([['ID'], [1]])
        response = self.enqueue('inspections/', content, content, 'input.xlsx',
                                {'selection': {'matching': {'mode': 'auto', 'unrelated': 'must-not-persist'}}})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(Job.objects.count(), 0)

    def test_excel_and_macro_workbooks_support_both_report_types(self):
        content = book([['ID', 'Value'], [1, '=1+1']])
        changed = book([['ID', 'Value'], [1, '=1+2']])
        for extension in ('xlsx', 'xlsm'):
            for output in ('word', 'excel'):
                with self.subTest(extension=extension, output=output):
                    source = self.run_job(self.enqueue('inspections/', content, changed, f'input.{extension}',
                                                       key=f'inspect-{extension}-{output}'))
                    response = self.client.post(BASE + 'jobs/', {'inspection_id': str(source.id),
                        'parameters': {'output_type': output}}, format='json',
                        HTTP_IDEMPOTENCY_KEY=f'compare-{extension}-{output}')
                    report = self.run_job(response)
                    self.assertEqual(report.result_summary['replace'], 1)
