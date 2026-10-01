from io import BytesIO
from django.test import SimpleTestCase
from openpyxl import Workbook

from comparison.contracts import CompareError, resolve_options
from comparison.excel import inspect_excel, compare_excel


def book(rows, title='Data'):
    workbook = Workbook()
    workbook.active.title = title
    for row in rows:
        workbook.active.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class ExcelMatchingTests(SimpleTestCase):
    def inspect(self, old, new, selection=None):
        return inspect_excel(old, new, selection or {}, resolve_options({}, 'excel'))

    def test_identifier_proposed_and_reordered_rows_match(self):
        old = book([['ID', 'Description'], [1, 'Alpha'], [2, 'Beta']])
        new = book([['ID', 'Description'], [2, 'Beta changed'], [1, 'Alpha']])
        inspected = self.inspect(old, new)
        self.assertFalse(inspected['matching']['requires_input'])
        self.assertEqual(inspected['matching']['method'], 'keys')
        result = compare_excel(old, new, inspected['selection'], resolve_options({}, 'excel'))
        self.assertEqual(result.counts['replace'], 1)
        changed = next(row for row in result.entries if row.change == 'replace')
        self.assertEqual((changed.old, changed.new), ('Beta', 'Beta changed'))
        self.assertIn('3', changed.old_location)
        self.assertIn('2', changed.new_location)

    def test_duplicate_candidates_require_guidance(self):
        old = book([['Name', 'State'], ['A', 'old'], ['A', 'old']])
        new = book([['Name', 'State'], ['A', 'new'], ['A', 'new']])
        self.assertTrue(self.inspect(old, new)['matching']['requires_input'])
        resolved = self.inspect(old, new, {'matching': {'mode': 'position'}})
        self.assertFalse(resolved['matching']['requires_input'])

    def test_duplicate_explicit_keys_rejected(self):
        rows = book([['ID', 'State'], [1, 'a'], [1, 'b']])
        with self.assertRaises(CompareError):
            self.inspect(rows, rows, {'matching': {'mode': 'keys', 'keys': [0]}})

    def test_added_columns_are_visible_even_when_existing_cells_unchanged(self):
        old = book([['ID', 'Name'], [1, 'Alpha']])
        new = book([['ID', 'Name', 'New field'], [1, 'Alpha', 'Value']])
        selection = self.inspect(old, new)['selection']
        result = compare_excel(old, new, selection, resolve_options({}, 'excel'))
        self.assertTrue(any(row.change == 'insert' and row.new == 'New field' for row in result.entries))
        self.assertTrue(any(row.change == 'insert' and row.new == 'Value' for row in result.entries))

    def test_formulas_are_compared_without_execution(self):
        old = book([['ID', 'Value'], [1, '=1+1']])
        new = book([['ID', 'Value'], [1, '=1+2']])
        result = compare_excel(old, new, self.inspect(old, new)['selection'], resolve_options({}, 'excel'))
        self.assertTrue(any(row.old == '=1+1' and row.new == '=1+2' for row in result.entries))

    def test_sheet_and_header_selection_with_renamed_column(self):
        old = book([['Title'], ['ID', 'Old name'], [1, 'a']], 'Before')
        new = book([['Title'], ['ID', 'New name'], [1, 'b']], 'After')
        selected = {'first': {'sheet': 'Before', 'header_row': 2},
                    'second': {'sheet': 'After', 'header_row': 2}, 'columns': [[0, 0], [1, 1]]}
        result = self.inspect(old, new, selected)
        self.assertEqual(result['first']['columns'], ['ID', 'Old name'])
        self.assertEqual(result['selection']['columns'], [[0, 0], [1, 1]])

    def test_auto_header_cannot_silently_exclude_populated_rows(self):
        old = book([['ID', 'Name', None], ['a', 'Original', None], ['b', 'Common', 'Notes'], ['c', 'Same', 'Value']])
        new = book([['ID', 'Name', None], ['a', 'Changed', None], ['b', 'Common', 'Notes'], ['c', 'Same', 'Value']])
        inspected = self.inspect(old, new)
        self.assertTrue(inspected['matching']['requires_input'])
        selection = {'first': {'sheet': 'Data', 'header_row': 1}, 'second': {'sheet': 'Data', 'header_row': 1}}
        result = compare_excel(old, new, selection, resolve_options({}, 'excel'))
        self.assertTrue(any(row.old == 'Original' and row.new == 'Changed' for row in result.entries))

    def test_unselected_sheet_does_not_expand_comparison_scope(self):
        workbook = Workbook()
        workbook.active.title = 'Selected'
        workbook.active.append(['ID', 'Name'])
        workbook.active.append([1, 'Value'])
        other = workbook.create_sheet('Unrelated')
        other.cell(20000, 1, 'Outside comparison scope')
        buffer = BytesIO()
        workbook.save(buffer)
        content = buffer.getvalue()
        selection = {'first': {'sheet': 'Selected', 'header_row': 1}, 'second': {'sheet': 'Selected', 'header_row': 1}}
        result = self.inspect(content, content, selection)
        self.assertFalse(result['matching']['requires_input'])

    def test_weak_threshold_also_controls_identity_matched_cell_strength(self):
        old = book([['ID', 'Value'], [1, 'abcde']])
        new = book([['ID', 'Value'], [1, 'abcdf']])
        selection = {'matching': {'mode': 'keys', 'keys': [0]}}
        for weak, expected in ((.7, 'possible'), (.85, 'none')):
            options = resolve_options({'preset': 'custom', 'equal_ratio': .9, 'weak_equal_ratio': weak}, 'excel')
            report = compare_excel(old, new, selection, options)
            changed = next(row for row in report.entries if row.change == 'replace')
            self.assertEqual(changed.strength, expected)
            self.assertEqual((changed.old, changed.new), ('abcde', 'abcdf'))
