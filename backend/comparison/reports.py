"""DOCX and XLSX renderers for the same complete comparison result."""

from difflib import SequenceMatcher
import re

from .contracts import check_limit

COLUMNS = ['Change', 'Match strength', 'Similarity', 'Field', 'Old location', 'New location', 'Old content', 'New content']


def metadata(result):
    return [('Format', result.family), ('Matching method', result.method), ('Preset', result.options.preset),
            ('Strong threshold', result.options.equal_ratio), ('Possible threshold', result.options.weak_equal_ratio),
            *[(name, value) for name, value in result.counts.items()],
            ('Scope', 'Text and cell content only. Similarity never hides a real difference.'),
            ('Count unit', 'Cells and column headers' if result.family == 'excel' else 'Text blocks'),
            *[('Warning', warning) for warning in result.warnings]]


def write_report(result, path, checkpoint=lambda: None):
    check_limit(len(result.entries) > 100000, 'The report exceeds 100,000 entries. Compare smaller tables or sections.')
    if result.options.output_type == 'excel':
        write_excel(result, path, checkpoint)
    else:
        check_limit(len(result.entries) > 10000, 'The Word report exceeds 10,000 entries. Select an Excel report.')
        write_word(result, path, checkpoint)


def safe_cell(value):
    from awcenter.spreadsheet_security import spreadsheet_safe_value
    if isinstance(value, str):
        check_limit(len(value) > 32760, 'A report cell exceeds the Excel text limit. Select a Word report.')
    return spreadsheet_safe_value(value)


def entry_values(entry):
    return [entry.change, entry.strength, entry.similarity, entry.field, entry.old_location,
            entry.new_location, entry.old, entry.new]


def write_excel(result, path, checkpoint):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Summary'
    for pair in metadata(result):
        summary.append([safe_cell(value) for value in pair])
    differences = workbook.create_sheet('Differences')
    possible = workbook.create_sheet('Possible matches')
    for sheet in (differences, possible):
        sheet.append(COLUMNS)
        sheet.freeze_panes = 'A2'
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in ('G', 'H'):
            sheet.column_dimensions[column].width = 70
        for column in ('A', 'B', 'D', 'E', 'F'):
            sheet.column_dimensions[column].width = 24
    for index, entry in enumerate(result.entries):
        if index % 200 == 0:
            checkpoint()
        values = [safe_cell(value) for value in entry_values(entry)]
        differences.append(values)
        if entry.strength == 'possible':
            possible.append(values)
    for sheet in (differences, possible):
        sheet.auto_filter.ref = sheet.dimensions
    workbook.save(path)
    workbook.close()


def write_word(result, path, checkpoint):
    from docx import Document

    document = Document()
    document.add_heading('Comparison report', 0)
    for key, value in metadata(result):
        document.add_paragraph(f'{key}: {value}')
    document.add_paragraph('Red strikethrough: removed content. Green underline: added content. '
                           'Possible matches require review; every real difference remains visible.')
    for index, entry in enumerate(result.entries):
        if index % 50 == 0:
            checkpoint()
        document.add_heading(f'{entry.change.title()} — {entry.field or "Text"}', 2)
        document.add_paragraph(f'{entry.old_location or "—"} → {entry.new_location or "—"} | {entry.strength}')
        if entry.change == 'equal':
            document.add_paragraph(entry.old)
        else:
            add_redline(document.add_paragraph(), entry.old, entry.new)
    document.save(path)


def add_redline(paragraph, old, new):
    from docx.shared import RGBColor

    first = re.findall(r'\s+|\w+|[^\w\s]', old)
    second = re.findall(r'\s+|\w+|[^\w\s]', new)
    for tag, start, end, next_start, next_end in SequenceMatcher(None, first, second, autojunk=False).get_opcodes():
        if tag == 'equal':
            paragraph.add_run(''.join(first[start:end]))
        if tag in ('delete', 'replace'):
            run = paragraph.add_run(''.join(first[start:end]))
            run.font.strike = True
            run.font.color.rgb = RGBColor.from_string('B42318')
        if tag in ('insert', 'replace'):
            run = paragraph.add_run(''.join(second[next_start:next_end]))
            run.font.underline = True
            run.font.color.rgb = RGBColor.from_string('067647')
