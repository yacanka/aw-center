"""Bounded OOXML table reading and explainable header suggestions."""

from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO
from collections import Counter

from openpyxl import load_workbook

from .contracts import CompareError, MAX_CHARACTERS, MAX_COLUMNS, MAX_ROWS, MAX_SHEETS, check_limit


@dataclass
class Table:
    sheet: str
    header_row: int
    columns: list[str]
    rows: list[tuple[int, list]]


def display(value):
    if value is None:
        return ''
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return str(value)


def read_tables(content, selection, checkpoint=lambda: None):
    if not isinstance(selection, dict):
        raise CompareError('Select a worksheet and header row.')
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        check_limit(len(workbook.worksheets) > MAX_SHEETS, 'The workbook has too many worksheets.')
        sheets = inspect_headers(workbook, checkpoint)
        if not sheets:
            raise CompareError('No table with a header in the first 20 rows was found.', 'COMPARE_TABLE_EMPTY')
        name = selection.get('sheet', sheets[0]['name'])
        metadata = next((sheet for sheet in sheets if sheet['name'] == name), None)
        if metadata is None:
            raise CompareError('Select a worksheet present in this workbook.')
        rows = read_selected_sheet(workbook[name], checkpoint)
        header = selection.get('header_row', metadata['headers'][0]['row'])
        if type(header) is not int or not 1 <= header <= min(20, len(rows)):
            raise CompareError('Select a header row from the first 20 rows.')
        labels = [display(value) for value in rows[header - 1]]
        if not any(labels):
            raise CompareError('The selected header row is empty.')
        columns = [label or f'Column {index + 1}' for index, label in enumerate(labels)]
        data = [(index, row) for index, row in enumerate(rows[header:], header + 1)
                if any(value is not None for value in row)]
        check_limit(len(data) > MAX_ROWS, 'The table exceeds the 10,000-row limit.')
        table = Table(name, header, columns, data)
        excluded = sum(any(value is not None for value in row) for row in rows[:header - 1])
        info = {'sheets': sheets, 'selected': {'sheet': name, 'header_row': header},
                'columns': columns, 'preview': [[display(v) for v in row] for _, row in data[:5]],
                'row_count': len(data), 'excluded_rows': excluded,
                'header_confirmation_required': excluded > 0 and 'header_row' not in selection}
        return table, info
    finally:
        workbook.close()


def inspect_headers(workbook, checkpoint):
    sheets = []
    characters = 0
    for sheet in workbook.worksheets:
        checkpoint()
        # Sampling other sheets must not pull their full content into the
        # selected-table comparison or allocate their declared dimensions.
        width = min(sheet.max_column or MAX_COLUMNS, MAX_COLUMNS)
        rows = list(sheet.iter_rows(max_row=20, max_col=width, values_only=True))
        characters += sum(len(display(value)) for row in rows for value in row)
        check_limit(characters > MAX_CHARACTERS, 'The worksheet previews exceed the extracted content limit.')
        candidates = header_candidates(rows)
        if candidates:
            sheets.append({'name': sheet.title, 'headers': candidates})
    return sheets


def read_selected_sheet(sheet, checkpoint):
    check_limit(sheet.max_column and sheet.max_column > MAX_COLUMNS, 'The selected worksheet exceeds the 100-column limit.')
    check_limit(sheet.max_row and sheet.max_row > MAX_ROWS + 20, 'The selected worksheet exceeds the 10,000-row table limit.')
    rows = []
    characters = 0
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
        check_limit(row_number > MAX_ROWS + 20 or len(row) > MAX_COLUMNS, 'The selected worksheet exceeds comparison limits.')
        characters += sum(len(display(value)) for value in row)
        check_limit(characters > MAX_CHARACTERS, 'The selected worksheet exceeds the extracted content limit.')
        rows.append(list(row))
        if row_number % 250 == 0:
            checkpoint()
    while rows and not any(value is not None for value in rows[-1]):
        rows.pop()
    width = max((index + 1 for row in rows for index, value in enumerate(row) if value is not None), default=0)
    return [(row + [None] * width)[:width] for row in rows]


def header_candidates(rows):
    candidates = []
    for number, row in enumerate(rows[:20], 1):
        values = [display(value).strip() for value in row if value is not None]
        if not values:
            continue
        text_count = sum(isinstance(value, str) and not value.startswith('=') for value in row)
        score = (text_count, len(set(values)), -number)
        candidates.append((score, {'row': number, 'labels': [display(v)[:120] for v in row]}))
    return [item for _, item in sorted(candidates, key=lambda item: item[0], reverse=True)]


def column_pairs(old, new, selected=None):
    if selected is None:
        left = Counter(label.strip().casefold() for label in old.columns)
        right = Counter(label.strip().casefold() for label in new.columns)
        lookup = {label.strip().casefold(): index for index, label in enumerate(new.columns)}
        return [[index, lookup[label.strip().casefold()]] for index, label in enumerate(old.columns)
                if left[label.strip().casefold()] == right[label.strip().casefold()] == 1]
    if not isinstance(selected, list) or len(selected) > MAX_COLUMNS:
        raise CompareError('Select valid column mappings.')
    used_old, used_new = set(), set()
    for pair in selected:
        if (not isinstance(pair, list) or len(pair) != 2
                or any(type(index) is not int for index in pair)):
            raise CompareError('Select valid column mappings.')
        first, second = pair
        if not 0 <= first < len(old.columns) or not 0 <= second < len(new.columns):
            raise CompareError('A selected column does not exist.')
        if first in used_old or second in used_new:
            raise CompareError('Each column may be mapped only once.')
        used_old.add(first)
        used_new.add(second)
    return selected
