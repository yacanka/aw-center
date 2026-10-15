"""Preview/confirm import matching and concurrency acceptance tests."""

import json
import re

from datetime import date
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile

import pandas as pd
from openpyxl import Workbook
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.db.models.signals import post_init
from django.test import TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from rest_framework.test import APIClient
from rest_framework.exceptions import ValidationError

from orgs.models import Panel, Person, Project, ProjectRoleAssignment, ResponsibleAssignment

from .models import ComplianceDocument, CoverPage, ImportAudit, TrackingProfile, WorkflowEvent
from .imports import execute_plan, plan_fingerprint, prepare_plan, prepare_tabular_plan


class ComplianceImportTests(TestCase):
    def setUp(self):
        self.project = Project.objects.get(slug="ozgur")
        self.user = get_user_model().objects.create_user("import-editor")
        ProjectRoleAssignment.objects.create(
            project=self.project,
            domain=ProjectRoleAssignment.Domain.COMPLIANCE,
            role=ProjectRoleAssignment.Role.EDITOR,
            user=self.user,
        )
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.panel = Panel.objects.create(
            project=self.project,
            name="Flight Controls",
            ata="27-00",
        )

    def workbook(
        self,
        *,
        name="Imported Document",
        cover="CP-I",
        tech="TD-I",
        status=None,
        effective_date=None,
        target_date=None,
        revised_target_date=None,
        delivery_date=None,
        panel=None,
        ata=None,
        moc=None,
    ):
        output = BytesIO()
        row = {
            "Document Name": name,
            "Cover Page Number": cover,
            "Technical Document No": tech,
        }
        if status is not None:
            row["Status"] = status
        if effective_date is not None:
            row["Effective Date"] = effective_date
        if target_date is not None:
            row["UBM Target Date"] = target_date
        if revised_target_date is not None:
            row["UBM Revised Target Date"] = revised_target_date
        if delivery_date is not None:
            row["UBM Delivery Date"] = delivery_date
        if panel is not None:
            row["Panel"] = panel
        if ata is not None:
            row["ATA Chapter"] = ata
        if moc is not None:
            row["MOC"] = moc
        pd.DataFrame([row]).to_excel(output, index=False)
        return output.getvalue()

    @staticmethod
    def upload(content):
        return SimpleUploadedFile(
            "compliance.xlsx",
            content,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def preview(self, content):
        return self.client.post(
            "/api/projects/ozgur/compliance-documents/imports/preview/",
            {"file": self.upload(content)},
            format="multipart",
        )

    def confirm(self, content, token):
        return self.client.post(
            "/api/projects/ozgur/compliance-documents/imports/confirm/",
            {"file": self.upload(content), "confirmation_token": token},
            format="multipart",
        )

    def manual_request(self, content, mapping, token=None):
        action = "confirm" if token is not None else "preview"
        data = {"file": self.upload(content), "mapping": json.dumps(mapping)}
        if token is not None:
            data["confirmation_token"] = token
        return self.client.post(
            f"/api/projects/ozgur/compliance-documents/imports/{action}/",
            data,
            format="multipart",
        )

    def test_terminal_audit_save_failure_rolls_back_all_import_mutations(self):
        for partial in (False, True):
            with self.subTest(partial=partial):
                output = BytesIO()
                rows = [{"name": f"Atomic document {partial}", "cover_page_no": f"CP-ATOMIC-{partial}",
                         "panel": "New panel", "ata": 29, "status": "Reviewed"}]
                if partial:
                    rows.append({"name": "Rejected document", "cover_page_no": None})
                pd.DataFrame(rows).to_excel(output, index=False)
                uploaded = self.upload(output.getvalue())
                request = SimpleNamespace(user=self.user)
                fingerprint = plan_fingerprint(prepare_plan(uploaded, self.project, request))
                original_save = ImportAudit.save

                def fail_terminal_save(audit, *args, **kwargs):
                    if audit.status in (ImportAudit.Status.SUCCESS, ImportAudit.Status.PARTIAL):
                        raise RuntimeError("Synthetic audit storage failure")
                    return original_save(audit, *args, **kwargs)

                with patch.object(ImportAudit, "save", fail_terminal_save):
                    with self.assertRaises(RuntimeError):
                        execute_plan(uploaded, self.project, request, fingerprint)

                self.assertFalse(ComplianceDocument.objects.exists())
                self.assertFalse(CoverPage.objects.filter(number=f"CP-ATOMIC-{partial}").exists())
                self.assertFalse(Panel.objects.filter(project=self.project, ata="29-00").exists())
                self.assertFalse(WorkflowEvent.objects.exists())
                audit = ImportAudit.objects.latest("started_at")
                self.assertEqual(audit.status, ImportAudit.Status.FAILED)
                self.assertIsNotNone(audit.completed_at)

    @override_settings(AWCENTER_MAX_COMPDOC_IMPORT_ROWS=2)
    def test_workbook_row_limit_bounds_reads_before_dataframe_allocation(self):
        output = BytesIO()
        pd.DataFrame([{"name": f"Document {index}", "cover_page_no": f"CP-{index}"}
                      for index in range(50)]).to_excel(output, index=False)
        original_read = pd.read_excel
        read_sizes = []

        def bounded_read(*args, **kwargs):
            frame = original_read(*args, **kwargs)
            read_sizes.append(len(frame))
            return frame

        with patch.object(pd, "read_excel", bounded_read):
            with self.assertRaises(ValidationError) as raised:
                prepare_plan(self.upload(output.getvalue()), self.project,
                             SimpleNamespace(user=self.user))
        self.assertEqual(raised.exception.get_codes()["file"], "IMPORT_ROW_LIMIT")
        self.assertLessEqual(max(read_sizes, default=0), 3)

    def test_sparse_wide_workbook_is_rejected_before_data_allocation(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["name", "cover_page_no"])
        for index in range(10):
            sheet.append([f"Document {index}", f"CP-{index}"])
        sheet["XFD2"] = "Sparse distant column"
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        for forged_dimension in (False, True):
            with self.subTest(forged_dimension=forged_dimension):
                content = output.getvalue()
                if forged_dimension:
                    rewritten = BytesIO()
                    with ZipFile(BytesIO(content)) as source:
                        with ZipFile(rewritten, "w") as target:
                            for entry in source.infolist():
                                data = source.read(entry.filename)
                                if entry.filename == "xl/worksheets/sheet1.xml":
                                    data = re.sub(rb'<dimension ref="[^"]+"',
                                                  b'<dimension ref="A1:B11"', data)
                                target.writestr(entry, data)
                    content = rewritten.getvalue()
                original_read = pd.read_excel
                body_reads = []

                def record_data_reads(*args, **kwargs):
                    if kwargs.get("nrows") != 0:
                        body_reads.append(True)
                    return original_read(*args, **kwargs)

                with patch.object(pd, "read_excel", record_data_reads):
                    with self.assertRaises(ValidationError) as raised:
                        prepare_plan(self.upload(content), self.project, SimpleNamespace(user=self.user))
                self.assertEqual(raised.exception.get_codes()["file"], "IMPORT_COLUMN_LIMIT")
                self.assertEqual(body_reads, [])

    def test_workbook_column_limit_accepts_boundary_with_unmapped_columns(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["name", "cover_page_no"] + [f"Extra {index}" for index in range(98)])
        sheet.append(["Boundary document", "CP-BOUNDARY"] + [None] * 97 + ["Extra value"])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        plan = prepare_plan(self.upload(output.getvalue()), self.project, SimpleNamespace(user=self.user))
        self.assertEqual(plan.counts["created_count"], 1)
        self.assertEqual(plan.errors, ())
        self.assertEqual(len(plan.mapping["source_columns"]), 100)
        self.assertEqual(len(plan.mapping["unmapped_columns"]), 98)

    @override_settings(AWCENTER_MAX_COMPDOC_IMPORT_ROWS=2)
    def test_sparse_far_row_is_rejected_instead_of_silently_truncated(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["name", "cover_page_no"])
        sheet["A24"] = "Distant document"
        sheet["B24"] = "CP-DISTANT"
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        for forged_dimension in (False, True):
            with self.subTest(forged_dimension=forged_dimension):
                content = output.getvalue()
                if forged_dimension:
                    rewritten = BytesIO()
                    with ZipFile(BytesIO(content)) as source:
                        with ZipFile(rewritten, "w") as target:
                            for entry in source.infolist():
                                data = source.read(entry.filename)
                                if entry.filename == "xl/worksheets/sheet1.xml":
                                    data = re.sub(rb'<dimension ref="[^"]+"',
                                                  b'<dimension ref="A1:B1"', data)
                                target.writestr(entry, data)
                    content = rewritten.getvalue()
                original_read = pd.read_excel
                body_reads = []

                def record_data_reads(*args, **kwargs):
                    if kwargs.get("nrows") != 0:
                        body_reads.append(True)
                    return original_read(*args, **kwargs)

                with patch.object(pd, "read_excel", record_data_reads):
                    with self.assertRaises(ValidationError) as raised:
                        prepare_plan(self.upload(content), self.project, SimpleNamespace(user=self.user))
                self.assertEqual(raised.exception.get_codes()["file"], "IMPORT_ROW_LIMIT")
                self.assertEqual(body_reads, [])

    @override_settings(AWCENTER_MAX_COMPDOC_IMPORT_ROWS=2)
    def test_row_limit_accepts_full_boundary_with_both_supported_header_offsets(self):
        for preamble in (False, True):
            with self.subTest(preamble=preamble):
                workbook = Workbook()
                sheet = workbook.active
                if preamble:
                    sheet.append(["AW Center export"])
                sheet.append(["name", "cover_page_no"])
                sheet.append(["First", "CP-1"])
                sheet.append(["Second", "CP-2"])
                output = BytesIO()
                workbook.save(output)
                workbook.close()
                plan = prepare_plan(self.upload(output.getvalue()), self.project,
                                    SimpleNamespace(user=self.user))
                self.assertEqual(plan.counts["created_count"], 2)
                self.assertEqual(plan.errors, ())
                self.assertEqual(plan.mapping["header_row"], 2 if preamble else 1)

    def test_preview_identity_validation_queries_do_not_grow_per_row(self):
        request = SimpleNamespace(user=self.user)

        def plan_rows(count):
            return tuple((index + 2, {"name": f"Doc {index}", "cover_page_no": f"CP-{index}"})
                         for index in range(count))

        with CaptureQueriesContext(connection) as one:
            first = prepare_tabular_plan(plan_rows(1), self.project, request, mapping={})
        with CaptureQueriesContext(connection) as many:
            second = prepare_tabular_plan(plan_rows(20), self.project, request, mapping={})
        self.assertEqual(first.counts["created_count"], 1)
        self.assertEqual(second.counts["created_count"], 20)
        self.assertLessEqual(len(many), len(one) + 1)

    def test_preview_does_not_hydrate_unrelated_document_payloads(self):
        cover = CoverPage.objects.create(project=self.project, number="UNRELATED")
        for index in range(15):
            ComplianceDocument.objects.create(project=self.project, cover_page=cover,
                                              name=f"Unrelated {index}", notes="Long payload " * 100)
        hydrated = []

        def record_document(sender, instance, **kwargs):
            hydrated.append(instance.pk)

        post_init.connect(record_document, sender=ComplianceDocument)
        try:
            plan = prepare_tabular_plan(((2, {"name": "New document", "cover_page_no": "NEW"}),),
                                       self.project, SimpleNamespace(user=self.user), mapping={})
        finally:
            post_init.disconnect(record_document, sender=ComplianceDocument)
        self.assertEqual(plan.counts["created_count"], 1)
        self.assertEqual(hydrated, [])

    def test_confirmation_parses_once_before_lock_and_rechecks_changed_target(self):
        cover = CoverPage.objects.create(project=self.project, number="CP-I")
        document = ComplianceDocument.objects.create(project=self.project, cover_page=cover,
                                                      name="Imported Document")
        content = self.workbook(tech="New technical document")
        preview = self.preview(content)
        original_read = pd.read_excel
        locked = False
        data_reads = []

        def observe_project_lock(execute, sql, params, many, context):
            nonlocal locked
            if sql.lstrip().upper().startswith('UPDATE "ORGS_PROJECT"'):
                locked = True
            return execute(sql, params, many, context)

        def read_then_change_target(*args, **kwargs):
            frame = original_read(*args, **kwargs)
            if kwargs.get("nrows") != 0:
                data_reads.append(locked)
                # Simulate an edit completing after source parsing and before
                # confirmation obtains the project fence.
                ComplianceDocument.objects.filter(pk=document.pk).update(
                    name="Concurrent edit", version=document.version + 1,
                )
            return frame

        with connection.execute_wrapper(observe_project_lock):
            with patch.object(pd, "read_excel", read_then_change_target):
                response = self.confirm(content, preview.data["confirmation_token"])
        self.assertEqual(data_reads, [False])
        self.assertTrue(locked)
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.data["code"], "VERSION_CONFLICT")
        document.refresh_from_db()
        self.assertEqual(document.name, "Concurrent edit")
        self.assertIsNone(document.number)
        self.assertEqual(ImportAudit.objects.get().status, ImportAudit.Status.FAILED)

    def test_confirmation_source_rejection_still_records_failed_audit(self):
        output = BytesIO()
        pd.DataFrame([{"name": "First", "cover_page_no": "CP-1"},
                      {"name": "Second", "cover_page_no": "CP-2"}]).to_excel(output, index=False)
        uploaded = self.upload(output.getvalue())
        request = SimpleNamespace(user=self.user)
        fingerprint = plan_fingerprint(prepare_plan(uploaded, self.project, request))
        with override_settings(AWCENTER_MAX_COMPDOC_IMPORT_ROWS=1):
            with self.assertRaises(ValidationError):
                execute_plan(uploaded, self.project, request, fingerprint)
        self.assertEqual(ImportAudit.objects.count(), 1)
        audit = ImportAudit.objects.get()
        self.assertEqual(audit.status, ImportAudit.Status.FAILED)
        self.assertIsNotNone(audit.completed_at)
        self.assertFalse(ComplianceDocument.objects.exists())

    def test_candidate_selection_preserves_unicode_casefold_and_project_scope(self):
        cover = CoverPage.objects.create(project=self.project, number="CP-STRASSE")
        document = ComplianceDocument.objects.create(project=self.project, cover_page=cover,
                                                      name="Straße")
        other_project = Project.objects.exclude(pk=self.project.pk).first()
        other_cover = CoverPage.objects.create(project=other_project, number="CP-STRASSE")
        other_document = ComplianceDocument.objects.create(project=other_project,
                                                           cover_page=other_cover, name="STRASSE")
        plan = prepare_tabular_plan(((2, {"name": "STRASSE", "cover_page_no": "cp-straße"}),),
                                   self.project, SimpleNamespace(user=self.user), mapping={})
        self.assertEqual(plan.errors, ())
        self.assertEqual(plan.rows[0].target.pk, document.pk)
        self.assertNotEqual(plan.rows[0].target.pk, other_document.pk)
        foreign_id_plan = prepare_tabular_plan(
            ((2, {"id": str(other_document.pk), "name": "STRASSE", "cover_page_no": "cp-straße"}),),
            self.project, SimpleNamespace(user=self.user), mapping={},
        )
        self.assertEqual(foreign_id_plan.rows, ())
        self.assertEqual(foreign_id_plan.errors[0]["code"], "IMPORT_DOCUMENT_NOT_FOUND")

    def test_import_normalizes_single_and_multiple_mocs(self):
        cases = [
            ("0", "0"), ("9", "9"), (" M ", "M"), (0, "0"), (9, "9"),
            (2.0, "2"), (" 2 ", "2"),
            ("1,2", "M"), ("1;2", "M"), ("1 2", "M"),
            ("1\t2", "M"), ("1\n2", "M"), ("1\r\n2", "M"),
            ("1/2", "M"), ("1|2", "M"), ("1-2", "M"),
            (" 1 / 2;\n3 ", "M"), ("1,,  2", "M"), ("1,1", "M"),
        ]
        for value, expected in cases:
            with self.subTest(moc=value):
                plan = prepare_tabular_plan(
                    ((2, {"name": "MOC document", "cover_page_no": "CP-M", "moc": value}),),
                    self.project, SimpleNamespace(user=self.user), mapping={},
                )
                self.assertEqual(plan.errors, ())
                self.assertEqual(plan.rows[0].payload["moc"], expected)

    def test_import_rejects_invalid_mocs_with_field_errors(self):
        for value in ("12", "1.2", "2.0", "A", "1,A", "m", "1,", ",1", "1:M", "1:2", "M,1", 12, 1.2, True):
            with self.subTest(moc=value):
                plan = prepare_tabular_plan(
                    ((2, {"name": "Invalid MOC", "cover_page_no": "CP-M", "moc": value}),),
                    self.project, SimpleNamespace(user=self.user), mapping={},
                )
                self.assertEqual(plan.counts["rejected_count"], 1)
                self.assertIn("moc", plan.errors[0]["fields"])
                self.assertEqual(plan.rows, ())

    def test_excel_moc_preview_confirm_and_reimport(self):
        initial = self.workbook(moc=2.0)
        preview = self.preview(initial)
        self.assertEqual(preview.data["rejected_count"], 0, preview.data)
        self.assertEqual(self.confirm(initial, preview.data["confirmation_token"]).status_code, 201)
        document = ComplianceDocument.objects.get()
        self.assertEqual(document.moc, "2")

        multiple = self.workbook(moc="1,2\n3")
        preview = self.preview(multiple)
        self.assertEqual(preview.data["updated_count"], 1, preview.data)
        self.assertEqual(self.confirm(multiple, preview.data["confirmation_token"]).status_code, 201)
        document.refresh_from_db()
        self.assertEqual(document.moc, "M")
        version = document.version
        history_count = document.history.count()

        for content in (multiple, self.workbook(moc=""), self.workbook()):
            with self.subTest(content_size=len(content)):
                preview = self.preview(content)
                self.assertEqual(preview.data["unchanged_count"], 1, preview.data)
                self.assertEqual(self.confirm(content, preview.data["confirmation_token"]).status_code, 201)
                document.refresh_from_db()
                self.assertEqual(document.moc, "M")
                self.assertEqual(document.version, version)
                self.assertEqual(document.history.count(), history_count)

    def test_excel_multiple_mocs_are_saved_on_creation(self):
        content = self.workbook(moc="0 / 9")
        preview = self.preview(content)
        self.assertEqual(preview.data["created_count"], 1, preview.data)
        self.assertEqual(self.confirm(content, preview.data["confirmation_token"]).status_code, 201)
        self.assertEqual(ComplianceDocument.objects.get().moc, "M")

    def test_manual_mapping_recovers_unrecognized_headers_and_audits_links(self):
        output = BytesIO()
        pd.DataFrame([{"Custom heading": "Manual document", "Custom notes": "Note", "Custom cover": "CP-M"}]).to_excel(output, index=False)
        content = output.getvalue()
        automatic = self.preview(content)
        self.assertEqual(automatic.status_code, 200)
        self.assertEqual(automatic.data["missing_columns"], ["name", "cover_page_no"])
        self.assertEqual(automatic.data["confirmation_token"], "")
        self.assertEqual(automatic.data["source_columns"], ["Custom heading", "Custom notes", "Custom cover"])
        mapping = {"Custom heading": "name", "Custom notes": "notes", "Custom cover": "cover_page_no"}
        preview = self.manual_request(content, mapping)
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.data["created_count"], 1)
        self.assertFalse(ComplianceDocument.objects.exists())
        confirmed = self.manual_request(content, mapping, preview.data["confirmation_token"])
        self.assertEqual(confirmed.status_code, 201)
        self.assertEqual(ComplianceDocument.objects.get().name, "Manual document")
        self.assertEqual(ImportAudit.objects.get().mapped_columns, preview.data["mapped_columns"])

    def test_manual_confirmation_rejects_changed_or_omitted_links(self):
        content = self.workbook()
        mapping = {"Document Name": "name", "Cover Page Number": "cover_page_no"}
        preview = self.manual_request(content, mapping)
        token = preview.data["confirmation_token"]
        changed = self.manual_request(content, {"Technical Document No": "name", "Cover Page Number": "cover_page_no"}, token)
        self.assertEqual(changed.status_code, 400)
        self.assertEqual(self.confirm(content, token).status_code, 400)
        self.assertFalse(ComplianceDocument.objects.exists())

    def test_manual_mapping_rejects_invalid_links_and_requires_name(self):
        content = self.workbook()
        for mapping in [[], {"Document Name": "project"}, {"Missing": "name"},
                        {"Document Name": "name", "Technical Document No": "name"}]:
            with self.subTest(mapping=mapping):
                self.assertEqual(self.manual_request(content, mapping).status_code, 400)
        missing = self.manual_request(content, {})
        self.assertEqual(missing.status_code, 200)
        self.assertEqual(missing.data["confirmation_token"], "")
        self.assertFalse(ComplianceDocument.objects.exists())

    def test_manual_mapping_ignores_unselected_column_with_target_field_name(self):
        output = BytesIO()
        pd.DataFrame([{"Title": "Selected", "name": "Ignored", "Cover": "CP-M"}]).to_excel(output, index=False)
        content = output.getvalue()
        mapping = {"Title": "name", "Cover": "cover_page_no"}
        preview = self.manual_request(content, mapping)
        self.assertEqual(preview.data["created_count"], 1)
        self.assertEqual(self.manual_request(content, mapping, preview.data["confirmation_token"]).status_code, 201)
        self.assertEqual(ComplianceDocument.objects.get().name, "Selected")

    def test_import_without_cover_column_requires_a_new_preview(self):
        output = BytesIO()
        pd.DataFrame([{"Document Name": "Unnumbered", "Technical Document No": "TD-NEW"}]).to_excel(output, index=False)
        content = output.getvalue()
        preview = self.preview(content)
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.data["missing_columns"], ["cover_page_no"])
        self.assertEqual(preview.data["confirmation_token"], "")
        self.assertEqual(self.confirm(content, "").status_code, 400)
        self.assertFalse(ComplianceDocument.objects.exists())

    def test_blank_covers_are_rejected_without_creating_records(self):
        output = BytesIO()
        pd.DataFrame([
            {"Document Name": "First", "Cover Page Number": "", "Cover Page Issue": "A"},
            {"Document Name": "Second", "Cover Page Number": "", "Cover Page Issue": "B"},
        ]).to_excel(output, index=False)
        content = output.getvalue()
        preview = self.preview(content)
        self.assertEqual(preview.data["created_count"], 0, preview.data)
        self.assertEqual(preview.data["rejected_count"], 2)
        self.assertEqual(self.confirm(content, preview.data["confirmation_token"]).status_code, 201)
        self.assertFalse(CoverPage.objects.exists())
        self.assertFalse(ComplianceDocument.objects.exists())

    def test_import_by_uuid_still_requires_number_and_allows_explicit_rename(self):
        cover = CoverPage.objects.create(project=self.project, number="CP-ALLOCATED")
        document = ComplianceDocument.objects.create(project=self.project, cover_page=cover, name="Before")
        row = {"id": str(document.pk), "Document Name": "After"}
        output = BytesIO()
        pd.DataFrame([row]).to_excel(output, index=False)
        preview = self.preview(output.getvalue())
        self.assertEqual(preview.data["missing_columns"], ["cover_page_no"])
        self.assertEqual(preview.data["confirmation_token"], "")
        document.refresh_from_db()
        self.assertEqual(document.name, "Before")
        output = BytesIO()
        pd.DataFrame([{**row, "Cover Page Number": cover.number}]).to_excel(output, index=False)
        content = output.getvalue()
        preview = self.preview(content)
        self.assertEqual(preview.data["updated_count"], 1, preview.data)
        self.assertEqual(self.confirm(content, preview.data["confirmation_token"]).status_code, 201)
        document.refresh_from_db()
        self.assertEqual(document.name, "After")
        self.assertEqual(document.cover_page_id, cover.pk)
        self.assertEqual(document.cover_page.number, "CP-ALLOCATED")

    def test_preview_then_confirm_creates_document_and_audit(self):
        content = self.workbook()
        preview = self.preview(content)
        confirmed = self.confirm(content, preview.data["confirmation_token"])

        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.data["created_count"], 1)
        self.assertEqual(confirmed.status_code, 201)
        self.assertEqual(ComplianceDocument.objects.count(), 1)
        self.assertEqual(ImportAudit.objects.get().status, ImportAudit.Status.SUCCESS)

    def test_confirm_updates_existing_document_without_optional_panel(self):
        cover = CoverPage.objects.create(project=self.project, number="CP-I")
        document = ComplianceDocument.objects.create(
            project=self.project,
            cover_page=cover,
            name="Imported Document",
            number=None,
            panel=None,
        )
        content = self.workbook()

        preview = self.preview(content)
        confirmed = self.confirm(content, preview.data["confirmation_token"])

        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.data["updated_count"], 1)
        self.assertEqual(confirmed.status_code, 201)
        document.refresh_from_db()
        self.assertEqual(document.name, "Imported Document")
        self.assertIsNone(document.panel_id)

    def test_new_matching_target_after_preview_returns_version_conflict(self):
        content = self.workbook()
        preview = self.preview(content)
        cover = CoverPage.objects.create(project=self.project, number="CP-I")
        ComplianceDocument.objects.create(
            project=self.project,
            cover_page=cover,
            name="Imported Document",
            number="TD-I",
        )

        response = self.confirm(content, preview.data["confirmation_token"])

        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.data["code"], "VERSION_CONFLICT")
        audit = ImportAudit.objects.get()
        self.assertEqual(audit.status, ImportAudit.Status.FAILED)
        self.assertEqual(audit.error_summary[0]["code"], "VERSION_CONFLICT")

    def test_import_matches_existing_document_by_name_when_technical_number_is_added(self):
        cover = CoverPage.objects.create(project=self.project, number="CP-I")
        document = ComplianceDocument.objects.create(
            project=self.project,
            cover_page=cover,
            name="Imported Document",
            number=None,
        )
        content = self.workbook()

        preview = self.preview(content)
        confirmed = self.confirm(content, preview.data["confirmation_token"])

        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.data["updated_count"], 1)
        self.assertEqual(confirmed.status_code, 201)
        self.assertEqual(ComplianceDocument.objects.count(), 1)
        document.refresh_from_db()
        self.assertEqual(document.number, "TD-I")

    def test_workbook_change_after_preview_is_rejected(self):
        first = self.workbook(name="First")
        preview = self.preview(first)

        response = self.confirm(self.workbook(name="Second"), preview.data["confirmation_token"])

        self.assertEqual(response.status_code, 400)
        self.assertIn(response.data["code"], {"IMPORT_PREVIEW_MISMATCH", "VALIDATION_ERROR"})

    def test_preview_rejects_duplicate_explicit_document_ids(self):
        cover = CoverPage.objects.create(project=self.project, number="CP-DUP")
        document = ComplianceDocument.objects.create(
            project=self.project,
            cover_page=cover,
            name="Existing document",
            number="TD-DUP",
        )
        output = BytesIO()
        pd.DataFrame(
            [
                {
                    "id": str(document.pk),
                    "name": "First update",
                    "cover_page_no": "CP-DUP",
                },
                {
                    "id": str(document.pk),
                    "name": "Second update",
                    "cover_page_no": "CP-DUP",
                },
            ]
        ).to_excel(output, index=False)

        preview = self.preview(output.getvalue())

        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.data["updated_count"], 0)
        self.assertEqual(preview.data["rejected_count"], 2)
        self.assertEqual(
            {item["code"] for item in preview.data["invalid_documents"]},
            {"IMPORT_DUPLICATE_DOCUMENT_ID"},
        )

    def test_preview_rejects_conflicting_issues_for_a_shared_cover_page(self):
        output = BytesIO()
        pd.DataFrame(
            [
                {
                    "name": "First shared-cover document",
                    "cover_page_no": "CP-SHARED",
                    "cover_page_issue": "A",
                },
                {
                    "name": "Second shared-cover document",
                    "cover_page_no": "CP-SHARED",
                    "cover_page_issue": "B",
                },
            ]
        ).to_excel(output, index=False)

        preview = self.preview(output.getvalue())

        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.data["created_count"], 0)
        self.assertEqual(preview.data["rejected_count"], 2)
        self.assertEqual(
            {item["code"] for item in preview.data["invalid_documents"]},
            {"IMPORT_CONFLICTING_COVER_PAGE_ISSUE"},
        )

    def test_import_canonicalizes_uppercase_explicit_document_id(self):
        cover = CoverPage.objects.create(project=self.project, number="CP-ID")
        document = ComplianceDocument.objects.create(
            project=self.project,
            cover_page=cover,
            name="Before ID import",
        )
        output = BytesIO()
        pd.DataFrame(
            [
                {
                    "id": str(document.pk).upper(),
                    "name": "After ID import",
                    "cover_page_no": cover.number,
                }
            ]
        ).to_excel(output, index=False)
        content = output.getvalue()

        preview = self.preview(content)
        confirmed = self.confirm(content, preview.data["confirmation_token"])

        self.assertEqual(preview.data["updated_count"], 1)
        self.assertEqual(confirmed.status_code, 201)
        document.refresh_from_db()
        self.assertEqual(document.name, "After ID import")

    def test_import_resolves_flexibly_formatted_ata_column(self):
        for index, ata_value in enumerate((27, 2700, "27-00", "27.00", "ATA 27"), start=1):
            with self.subTest(ata_value=ata_value):
                content = self.workbook(
                    name=f"ATA document {index}",
                    cover=f"CP-ATA-{index}",
                    tech=f"TD-ATA-{index}",
                    ata=ata_value,
                )

                preview = self.preview(content)
                confirmed = self.confirm(content, preview.data["confirmation_token"])

                self.assertEqual(preview.status_code, 200)
                self.assertIn(
                    {"source": "ATA Chapter", "target": "ata"},
                    preview.data["mapped_columns"],
                )
                self.assertEqual(confirmed.status_code, 201)
                self.assertEqual(
                    ComplianceDocument.objects.get(name=f"ATA document {index}").panel,
                    self.panel,
                )

    def test_existing_panel_column_also_accepts_short_ata_format(self):
        content = self.workbook(panel=27)

        preview = self.preview(content)
        confirmed = self.confirm(content, preview.data["confirmation_token"])

        self.assertEqual(preview.status_code, 200)
        self.assertEqual(confirmed.status_code, 201)
        self.assertEqual(ComplianceDocument.objects.get().panel, self.panel)

    def test_ata_disambiguates_a_panel_name_with_multiple_chapters(self):
        Panel.objects.create(
            project=self.project,
            name=self.panel.name,
            ata="28-00",
        )
        content = self.workbook(panel=self.panel.name, ata=27)

        preview = self.preview(content)
        confirmed = self.confirm(content, preview.data["confirmation_token"])

        self.assertEqual(preview.status_code, 200)
        self.assertEqual(confirmed.status_code, 201)
        self.assertEqual(ComplianceDocument.objects.get().panel, self.panel)

    def test_existing_ata_updates_panel_name_and_preserves_identity(self):
        Panel.objects.create(
            project=self.project,
            name="Electrical",
            ata="24-00",
        )
        content = self.workbook(panel="Electrical", ata=27)

        preview = self.preview(content)

        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.data["rejected_count"], 0)
        self.panel.refresh_from_db()
        self.assertEqual(self.panel.name, "Flight Controls")
        confirmed = self.confirm(content, preview.data["confirmation_token"])
        self.assertEqual(confirmed.status_code, 201)
        self.panel.refresh_from_db()
        self.assertEqual(self.panel.name, "Electrical")
        self.assertEqual(ComplianceDocument.objects.get().panel_id, self.panel.pk)
        self.assertEqual(Panel.objects.get(project=self.project, ata="24-00").name, "Electrical")

    def test_new_panel_is_only_created_on_confirmation(self):
        content = self.workbook(panel="Hydraulics", ata=29)
        preview = self.preview(content)
        self.assertEqual(preview.data["rejected_count"], 0)
        self.assertFalse(Panel.objects.filter(project=self.project, ata="29-00").exists())
        self.assertEqual(preview.data["panel_changes"], [{
            "ata": "29-00", "old_name": None, "new_name": "Hydraulics", "action": "create",
        }])
        confirmed = self.confirm(content, preview.data["confirmation_token"])
        self.assertEqual(confirmed.status_code, 201)
        panel = Panel.objects.get(project=self.project, ata="29-00")
        self.assertEqual(ComplianceDocument.objects.get().panel, panel)
        repeated = self.preview(content)
        self.assertEqual(repeated.data["unchanged_count"], 1)
        self.assertEqual(repeated.data["panel_changes"], [])

    def test_last_valid_panel_name_wins_and_invalid_rows_have_no_side_effects(self):
        output = BytesIO()
        pd.DataFrame([
            {"Document Name": "First", "Cover Page Number": "CP-P", "Panel": "First name", "ATA Chapter": 29},
            {"Document Name": "Second", "Cover Page Number": "CP-P", "Panel": "Last valid name", "ATA Chapter": 29},
            {"Document Name": "Invalid", "Cover Page Number": "CP-P", "Panel": "Rejected name", "ATA Chapter": 29, "Notes": "x" * 5001},
            {"Document Name": "Invalid new", "Cover Page Number": "CP-P", "Panel": "Rejected new", "ATA Chapter": 30, "Notes": "x" * 5001},
        ]).to_excel(output, index=False)
        content = output.getvalue()
        preview = self.preview(content)
        self.assertEqual(preview.data["created_count"], 2)
        self.assertEqual(preview.data["rejected_count"], 2)
        response = self.confirm(content, preview.data["confirmation_token"])
        self.assertEqual(response.status_code, 201)
        panel = Panel.objects.get(project=self.project, ata="29-00")
        self.assertEqual(panel.name, "Last valid name")
        self.assertEqual(ComplianceDocument.objects.filter(panel=panel).count(), 2)
        self.assertFalse(Panel.objects.filter(project=self.project, ata="30-00").exists())

    def test_panel_change_after_preview_requires_new_confirmation(self):
        content = self.workbook(panel="Renamed", ata=27)
        preview = self.preview(content)
        self.panel.name = "Changed concurrently"
        self.panel.save()
        response = self.confirm(content, preview.data["confirmation_token"])
        self.assertEqual(response.status_code, 409)
        self.assertFalse(ComplianceDocument.objects.exists())
        self.panel.refresh_from_db()
        self.assertEqual(self.panel.name, "Changed concurrently")

    def test_new_panels_require_both_fields_and_valid_values(self):
        for panel, ata in [("New", None), (None, 29), ("New", "invalid 27"), ("x" * 256, 29)]:
            with self.subTest(panel=panel, ata=ata):
                content = self.workbook(panel=panel, ata=ata)
                preview = self.preview(content)
                self.assertEqual(preview.data["rejected_count"], 1)
                self.assertEqual(preview.data["panel_changes"], [])
                self.assertEqual(self.confirm(content, preview.data["confirmation_token"]).status_code, 201)
        self.assertEqual(Panel.objects.filter(project=self.project).count(), 1)
        self.assertFalse(ComplianceDocument.objects.exists())

    def test_digit_free_ata_text_is_preserved_through_preview_and_confirmation(self):
        content = self.workbook(panel="General panel", ata="  Not  applicable  ")
        preview = self.preview(content)
        self.assertEqual(preview.data["rejected_count"], 0, preview.data)
        self.assertEqual(preview.data["created_count"], 1)
        self.assertEqual(self.confirm(content, preview.data["confirmation_token"]).status_code, 201)
        document = ComplianceDocument.objects.get()
        self.assertEqual(document.panel.ata, "Not  applicable")
        self.assertEqual(document.panel.name, "General panel")

    def test_catalog_updates_are_project_scoped(self):
        other = Project.objects.exclude(pk=self.project.pk).first()
        panel = Panel.objects.create(project=other, ata="29-00", name="Other project")
        content = self.workbook(panel="Local", ata=29)
        preview = self.preview(content)
        self.assertEqual(self.confirm(content, preview.data["confirmation_token"]).status_code, 201)
        panel.refresh_from_db()
        self.assertEqual(panel.name, "Other project")
        self.assertEqual(ComplianceDocument.objects.get().panel.project_id, self.project.pk)

    def test_catalog_rename_applies_even_when_document_is_unchanged(self):
        content = self.workbook(panel="Flight Controls", ata=27)
        preview = self.preview(content)
        self.confirm(content, preview.data["confirmation_token"])
        document = ComplianceDocument.objects.get()
        version = document.version
        self.panel.discipline = "Preserved"
        self.panel.save()
        content = self.workbook(panel="Renamed", ata=27)
        preview = self.preview(content)
        self.assertEqual(preview.data["unchanged_count"], 1)
        self.assertEqual(preview.data["panel_changes"][0]["old_name"], "Flight Controls")
        self.assertEqual(self.confirm(content, preview.data["confirmation_token"]).status_code, 201)
        self.panel.refresh_from_db()
        document.refresh_from_db()
        self.assertEqual(self.panel.name, "Renamed")
        self.assertEqual(self.panel.discipline, "Preserved")
        self.assertEqual(document.version, version)
        self.assertEqual(document.panel_id, self.panel.pk)

    def test_new_panel_attaches_to_existing_unassigned_document(self):
        content = self.workbook()
        preview = self.preview(content)
        self.confirm(content, preview.data["confirmation_token"])
        content = self.workbook(panel="New", ata=29)
        preview = self.preview(content)
        self.assertEqual(preview.data["updated_count"], 1)
        self.assertEqual(self.confirm(content, preview.data["confirmation_token"]).status_code, 201)
        self.assertEqual(ComplianceDocument.objects.get().panel.ata, "29-00")

    def test_new_panel_inserted_after_preview_rejects_confirmation(self):
        content = self.workbook(panel="New", ata=29)
        preview = self.preview(content)
        existing = Panel.objects.create(project=self.project, ata="29-00", name="Concurrent")
        self.assertEqual(self.confirm(content, preview.data["confirmation_token"]).status_code, 409)
        existing.refresh_from_db()
        self.assertEqual(existing.name, "Concurrent")
        self.assertFalse(ComplianceDocument.objects.exists())

    def test_insert_between_revalidation_and_apply_conflicts_instead_of_overwriting(self):
        from .imports import prepare_plan

        content = self.workbook(panel="Imported", ata=29)
        preview = self.preview(content)

        def prepare_with_competing_insert(*args, **kwargs):
            plan = prepare_plan(*args, **kwargs)
            Panel.objects.create(project=self.project, ata="29-00", name="Competing")
            return plan

        with patch("compliance.imports.prepare_plan", side_effect=prepare_with_competing_insert):
            response = self.confirm(content, preview.data["confirmation_token"])
        self.assertEqual(response.status_code, 409)
        self.assertFalse(ComplianceDocument.objects.exists())
        self.assertEqual(ImportAudit.objects.get().status, ImportAudit.Status.FAILED)

    def test_document_failure_rolls_back_panel_creation_and_rename(self):
        for ata in (27, 29):
            with self.subTest(ata=ata):
                content = self.workbook(panel="New", ata=ata)
                preview = self.preview(content)
                with patch("compliance.imports.ComplianceDocumentSerializer.create", side_effect=ValidationError("Failure")):
                    response = self.confirm(content, preview.data["confirmation_token"])
                self.assertEqual(response.status_code, 400)
                self.panel.refresh_from_db()
                self.assertEqual(self.panel.name, "Flight Controls")
                self.assertFalse(Panel.objects.filter(project=self.project, ata="29-00").exists())
                self.assertFalse(ComplianceDocument.objects.exists())

    def test_compliance_import_permission_does_not_grant_direct_organization_writes(self):
        denied = self.client.patch(
            f"/api/projects/ozgur/organization/panels/{self.panel.pk}/",
            {"name": "Direct rename"}, format="json",
        )
        self.assertEqual(denied.status_code, 403)
        content = self.workbook(panel="Imported rename", ata=27)
        preview = self.preview(content)
        self.assertEqual(self.confirm(content, preview.data["confirmation_token"]).status_code, 201)
        self.panel.refresh_from_db()
        self.assertEqual(self.panel.name, "Imported rename")
        role = ProjectRoleAssignment.objects.get(user=self.user, project=self.project)
        role.role = ProjectRoleAssignment.Role.VIEWER
        role.save()
        self.assertEqual(self.preview(content).status_code, 403)
        self.assertEqual(self.confirm(content, preview.data["confirmation_token"]).status_code, 403)

    def test_tracking_recipients_survive_rename_but_prevent_new_panel_attachment(self):
        person = Person.objects.create(person_id="import-responsible", name="Responsible", email="responsible@example.test")
        assignment = ResponsibleAssignment.objects.create(panel=self.panel, person=person, responsibility_role="AS")
        content = self.workbook(panel="Flight Controls", ata=27)
        preview = self.preview(content)
        self.confirm(content, preview.data["confirmation_token"])
        document = ComplianceDocument.objects.get()
        profile = TrackingProfile.objects.create(
            document=document, responsible_mode=TrackingProfile.ResponsibleMode.CUSTOM,
            notification_enabled=True, notification_events=["overdue"],
        )
        profile.responsible_people.add(person)
        content = self.workbook(panel="Renamed", ata=27)
        preview = self.preview(content)
        self.assertEqual(self.confirm(content, preview.data["confirmation_token"]).status_code, 201)
        assignment.refresh_from_db()
        self.assertEqual(assignment.panel_id, self.panel.pk)
        self.assertEqual(list(profile.responsible_people.all()), [person])
        content = self.workbook(panel="New", ata=29)
        preview = self.preview(content)
        self.assertEqual(preview.data["rejected_count"], 1)
        self.assertEqual(preview.data["panel_changes"], [])
        self.assertEqual(self.confirm(content, preview.data["confirmation_token"]).status_code, 201)
        document.refresh_from_db()
        self.assertEqual(document.panel_id, self.panel.pk)
        self.assertFalse(Panel.objects.filter(project=self.project, ata="29-00").exists())

    def test_import_builds_to_be_issued_event_without_status_flow_column(self):
        content = self.workbook(
            status="To Be Issued",
            target_date=date(2026, 9, 10),
        )

        preview = self.preview(content)
        confirmed = self.confirm(content, preview.data["confirmation_token"])

        self.assertEqual(preview.status_code, 200)
        self.assertEqual(confirmed.status_code, 201)
        document = ComplianceDocument.objects.get()
        event = document.workflow_events.get()
        self.assertEqual(document.status, "to_be_issued")
        self.assertEqual(document.target_date, date(2026, 9, 10))
        self.assertIsNone(document.release_date)
        self.assertEqual(
            (event.status, event.effective_date),
            ("to_be_issued", timezone.localdate()),
        )

    def test_import_appends_current_status_with_release_date(self):
        initial = self.workbook(
            status="To Be Issued",
            target_date=date(2026, 9, 10),
        )
        initial_preview = self.preview(initial)
        self.confirm(initial, initial_preview.data["confirmation_token"])
        delivered = self.workbook(
            status="Authority Review",
            target_date=date(2026, 9, 10),
            delivery_date=date(2026, 9, 18),
        )

        preview = self.preview(delivered)
        confirmed = self.confirm(delivered, preview.data["confirmation_token"])

        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.data["updated_count"], 1)
        self.assertEqual(confirmed.status_code, 201)
        document = ComplianceDocument.objects.get()
        events = list(
            document.workflow_events.order_by("sequence").values_list("status", "effective_date")
        )
        self.assertEqual(
            events,
            [
                ("to_be_issued", timezone.localdate()),
                ("authority_review", timezone.localdate()),
            ],
        )
        self.assertEqual(document.status, "authority_review")
        self.assertEqual(document.release_date, date(2026, 9, 18))

    def test_dates_are_independent_of_each_other_and_workflow_history(self):
        content = self.workbook(
            status="Authority Review",
            target_date=date(2026, 9, 20),
            revised_target_date=date(2026, 9, 25),
            delivery_date=date(2026, 9, 18),
        )
        preview = self.preview(content)
        self.assertEqual(preview.data["rejected_count"], 0, preview.data)
        self.assertEqual(self.confirm(content, preview.data["confirmation_token"]).status_code, 201)
        document = ComplianceDocument.objects.get()
        self.assertEqual(document.target_date, date(2026, 9, 20))
        self.assertEqual(document.revised_target_date, date(2026, 9, 25))
        self.assertEqual(document.release_date, date(2026, 9, 18))
        self.assertEqual(document.workflow_events.count(), 1)

        updated = self.workbook(
            status="Authority Review", target_date=date(2026, 10, 1),
            revised_target_date=date(2026, 9, 15), delivery_date=date(2026, 9, 10),
        )
        preview = self.preview(updated)
        self.assertEqual(preview.data["updated_count"], 1, preview.data)
        self.assertEqual(self.confirm(updated, preview.data["confirmation_token"]).status_code, 201)
        document.refresh_from_db()
        self.assertEqual(document.target_date, date(2026, 10, 1))
        self.assertEqual(document.revised_target_date, date(2026, 9, 15))
        self.assertEqual(document.release_date, date(2026, 9, 10))
        self.assertEqual(document.workflow_events.count(), 1)

        cleared = self.workbook(status="Authority Review", revised_target_date="")
        preview = self.preview(cleared)
        self.assertEqual(preview.data["updated_count"], 1, preview.data)
        self.assertEqual(self.confirm(cleared, preview.data["confirmation_token"]).status_code, 201)
        document.refresh_from_db()
        self.assertIsNone(document.revised_target_date)
        self.assertEqual(document.target_date, date(2026, 10, 1))
        self.assertEqual(document.release_date, date(2026, 9, 10))

    def test_delivery_date_without_target_is_accepted(self):
        content = self.workbook(status="To Be Issued", delivery_date=date(2026, 9, 18))
        preview = self.preview(content)
        self.assertEqual(preview.data["rejected_count"], 0, preview.data)
        self.assertEqual(self.confirm(content, preview.data["confirmation_token"]).status_code, 201)
        document = ComplianceDocument.objects.get()
        self.assertEqual(document.status, "to_be_issued")
        self.assertIsNone(document.target_date)
        self.assertEqual(document.release_date, date(2026, 9, 18))

    def test_legacy_future_workflow_event_does_not_block_independent_dates(self):
        cover = CoverPage.objects.create(project=self.project, number="CP-I")
        document = ComplianceDocument.objects.create(
            project=self.project, cover_page=cover, name="Imported Document",
            status="to_be_issued", target_date=date(2028, 9, 20),
        )
        WorkflowEvent.objects.create(
            document=document, sequence=1, status="to_be_issued",
            effective_date=date(2028, 9, 20), source=WorkflowEvent.Source.IMPORT,
        )
        content = self.workbook(
            status="Authority Review", target_date=date(2028, 9, 20),
            delivery_date=date(2028, 9, 10),
        )

        preview = self.preview(content)
        self.assertEqual(preview.data["rejected_count"], 0, preview.data)
        confirmed = self.confirm(content, preview.data["confirmation_token"])

        self.assertEqual(confirmed.status_code, 201, confirmed.data)
        document.refresh_from_db()
        self.assertEqual(document.release_date, date(2028, 9, 10))
        self.assertEqual(document.workflow_events.count(), 2)

    def test_explicit_out_of_order_transition_is_reported_in_preview(self):
        cover = CoverPage.objects.create(project=self.project, number="CP-I")
        document = ComplianceDocument.objects.create(
            project=self.project, cover_page=cover, name="Imported Document",
            status="to_be_issued",
        )
        WorkflowEvent.objects.create(
            document=document, sequence=1, status="to_be_issued",
            effective_date=date(2028, 9, 20), source=WorkflowEvent.Source.IMPORT,
        )
        response = self.preview(self.workbook(
            status="Authority Review", effective_date=date(2028, 9, 10),
            target_date=date(2028, 9, 20), delivery_date=date(2028, 9, 1),
        ))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["rejected_count"], 1)
        self.assertIn("effective_date", response.data["invalid_documents"][0]["fields"])

    def test_date_only_import_preserves_existing_status(self):
        initial = self.workbook(status="Authority Review")
        preview = self.preview(initial)
        self.assertEqual(self.confirm(initial, preview.data["confirmation_token"]).status_code, 201)
        document = ComplianceDocument.objects.get()
        date_only = self.workbook(revised_target_date=date(2026, 9, 25))

        preview = self.preview(date_only)
        self.assertEqual(preview.data["updated_count"], 1, preview.data)
        self.assertEqual(self.confirm(date_only, preview.data["confirmation_token"]).status_code, 201)
        document.refresh_from_db()
        self.assertEqual(document.status, "authority_review")
        self.assertEqual(document.revised_target_date, date(2026, 9, 25))
        self.assertEqual(document.workflow_events.count(), 1)
        unchanged = self.preview(date_only)
        self.assertEqual(unchanged.data["unchanged_count"], 1, unchanged.data)
