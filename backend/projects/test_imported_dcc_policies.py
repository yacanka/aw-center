"""Regression coverage for the September DCC tuning import."""

from copy import deepcopy
from django.test import SimpleTestCase, override_settings
from dcc.document_job import build_render_context
from dcc.document_fields import main_issue_fields, panel_fields
from projects.registry import find_project_by_jira_component, get_project_definition
from unittest.mock import patch
from datetime import date


class ImportedDccPolicyTests(SimpleTestCase):
    def context(self, slug, panels):
        return build_render_context({"schema_version": 2, "project_slug": slug, "placeholders": {"Panels": panels}})

    def test_new_template_aliases_preserve_original_fields_and_source(self):
        panel, _, _ = panel_fields({"summary": "Flight Panel review", "assignee": {"displayName": "Primary Person"}, "customfield_45421": {"displayName": "Candidate Person"}})
        original = deepcopy(panel)
        context = self.context("hys", [panel])
        self.assertEqual(context["panels"][0]["as_name"], "Primary PERSON")
        self.assertEqual(context["panels"][0]["panel_name"], "Flight")
        self.assertEqual(context["Panels"][0]["Panel_AS_Name"], "Primary PERSON, Candidate PERSON")
        self.assertEqual(context["flight_manuals_as_name"], "Primary PERSON")
        self.assertEqual(panel, original)

    @override_settings(GOKBEY_SOFTWARE_AS_NAME="Configured Signatory")
    def test_variant_dedicated_panels_merge_without_mutating_the_snapshot(self):
        panels = [
            {"Panel_Name": "Protection Panel", "as_name": "Primary", "candidate_as_name": "Candidate", "Affected_Requirements": "REQ-1"},
            {"Panel_Name": "Protection Panel", "as_name": "Other", "Affected_Requirements": "REQ-2"},
            {"Panel_Name": "Software Panel", "Panel_AS_Name": "Original"},
            {"Panel_Name": "RFM Panel", "Panel_Updated_Time": "01.09.2026"},
            {"Panel_Name": "OSD Panel"},
        ]
        original = deepcopy(panels)
        for slug in ("gokbey_jandarma", "gokbey_sivil"):
            context = self.context(slug, panels)
            self.assertEqual(len(context["panels"]), 1)
            self.assertEqual(context["panels"][0]["as_name"], "Configured Signatory")
            self.assertEqual(context["protection_as_name"], "Primary, Candidate, Other")
            self.assertEqual(context["protection_affected_requirements"], "REQ-1\nREQ-2")
            self.assertEqual(context["rfm_update_time"], "01.09.2026")
        self.assertEqual(panels, original)

    def test_hys_and_ozgur_ica_and_missing_assignee(self):
        for slug in ("hys", "ozgur"):
            context = self.context(slug, [{"Panel_Name": "ica Panel"}])
            self.assertEqual(context["ica_as_name"], "")
            self.assertEqual(len(context["Panels"]), 1)

    def test_hurkus_placeholder_keeps_all_panels(self):
        panels = [{"Panel_Name": "Flight Panel", "Affected_Requirements": "REQ-1"}]
        self.assertEqual(self.context("hurkus", panels)["Panels"], panels)
        self.assertEqual(get_project_definition("hurkus").capabilities, ("dcc",))

    def test_jira_component_names_accept_existing_and_corrected_spelling(self):
        for component in ("OZGUR", "Özgür", "ÖZGÜR"):
            self.assertEqual(find_project_by_jira_component(component).slug, "ozgur")
        self.assertEqual(find_project_by_jira_component("GÖKBEY SİVİL").slug, "gokbey_sivil")
        self.assertEqual(find_project_by_jira_component("HÜRKUŞ 2").slug, "hurkus")
        self.assertEqual(find_project_by_jira_component("PIKU").slug, "piku")

    @patch("dcc.document_fields.timezone.localdate", return_value=date(2026, 9, 22))
    def test_document_date_does_not_replace_the_source_update_date(self, _today):
        fields = main_issue_fields({"updated": "2026-09-01T12:00:00+03:00"})
        self.assertEqual(fields["Update_Time"], "22.09.2026")
        self.assertEqual(fields["Source_Updated_Time"], "01.09.2026")

    def test_real_docx_render_escapes_xml_and_supports_both_panel_contracts(self):
        from pathlib import Path
        from tempfile import TemporaryDirectory
        from docx import Document
        from docxtpl import DocxTemplate
        from dcc.document_job import render_document

        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.docx"
            output = Path(directory) / "output.docx"
            document = Document()
            document.add_paragraph("{{ Design_Change_Title }}")
            document.add_paragraph("{% for panel in panels %}{{ panel.panel_name }}: {{ panel.as_name }}{% endfor %}")
            document.add_paragraph("{% for panel in Panels %}{{ panel.Panel_AS_Name }}{% endfor %}")
            document.save(template)
            snapshot = {"schema_version": 2, "project_slug": "hurkus", "placeholders": {
                "Design_Change_Title": "A & B < C > D", "Panels": [{"Panel_Name": "Flight Panel", "Panel_AS_Name": "Example & Co"}],
            }}
            render_document(DocxTemplate(template), snapshot, output)
            paragraphs = [paragraph.text for paragraph in Document(output).paragraphs]
            self.assertEqual(paragraphs, ["A & B < C > D", "Flight: Example & Co", "Example & Co"])
