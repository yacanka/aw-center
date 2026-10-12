import json
from io import BytesIO
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document
from openpyxl import load_workbook
from rest_framework.test import APITestCase


def word_upload(lines=("Example paragraph",)):
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    stream = BytesIO()
    document.save(stream)
    return SimpleUploadedFile("input.docx", stream.getvalue(), content_type="application/octet-stream")


class WordComparisonApiTests(APITestCase):
    def setUp(self):
        self.client.force_authenticate(get_user_model().objects.create_user("word-compare-user"))
        self.client.raise_request_exception = False

    def compare(self, parameters=None, *, raw=None, first=None, second=None):
        body = {"first": first or word_upload(), "second": second or word_upload()}
        if raw is not None:
            body["json"] = raw
        elif parameters is not None:
            body["json"] = json.dumps(parameters)
        return self.client.post("/api/tools/word/compare/", body, format="multipart")

    def test_missing_malformed_or_incomplete_parameters_return_400(self):
        for raw in (None, "{", "[]", "null", "{}"):
            with self.subTest(raw=raw):
                response = self.compare(raw=raw)
                self.assertEqual(response.status_code, 400)
                self.assertEqual(response.data["code"], "WORD_COMPARE_INVALID")

    def test_invalid_thresholds_and_output_are_rejected(self):
        for equal, weak, output in ((True, .7, "word"), (".9", .7, "word"), (float("nan"), .7, "word"),
                                    (.9, float("inf"), "word"), (.9, 0, "word"), (.5, .7, "word"),
                                    (1.1, .7, "word"), (.9, .7, "pdf")):
            with self.subTest(equal=equal, weak=weak, output=output):
                response = self.compare({"equal_ratio": equal, "weak_equal_ratio": weak, "output_type": output})
                self.assertEqual(response.status_code, 400)
                self.assertEqual(response.data["code"], "WORD_COMPARE_INVALID")

    def test_very_large_integer_threshold_is_input_error(self):
        response = self.compare({"equal_ratio": 10 ** 1000, "weak_equal_ratio": .7, "output_type": "word"})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data["code"], "WORD_COMPARE_INVALID")

    def test_malformed_office_xml_is_input_error(self):
        original = word_upload()
        output = BytesIO()
        with ZipFile(original) as source, ZipFile(output, "w", ZIP_DEFLATED) as destination:
            for entry in source.infolist():
                destination.writestr(entry, b"<broken" if entry.filename == "word/document.xml" else source.read(entry.filename))
        corrupted = SimpleUploadedFile("input.docx", output.getvalue(), content_type="application/octet-stream")
        response = self.compare({"equal_ratio": .9, "weak_equal_ratio": .7, "output_type": "word"}, first=corrupted)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data["code"], "WORD_COMPARE_INVALID")

    def test_null_auto_thresholds_keep_excel_report_contract(self):
        response = self.compare({"equal_ratio": None, "weak_equal_ratio": None, "output_type": "excel"})
        self.assertEqual(response.status_code, 200)
        report = load_workbook(BytesIO(response.content))
        self.assertIn("Diff", report.sheetnames)
        self.assertEqual(report["Diff"]["D2"].value, "Example paragraph")

    @patch("word.views.MAX_BLOCKS", 2, create=True)
    def test_document_block_budget_is_enforced(self):
        response = self.compare({"equal_ratio": .9, "weak_equal_ratio": .7, "output_type": "word"},
                                first=word_upload(("One", "Two", "Three")))
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data["code"], "WORD_COMPARE_LIMIT_EXCEEDED")

    @patch("word.views.MAX_CANDIDATES", 2, create=True)
    def test_pair_budget_is_enforced_before_alignment(self):
        response = self.compare({"equal_ratio": .9, "weak_equal_ratio": .7, "output_type": "word"},
                                first=word_upload(("One", "Two", "Three")))
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data["code"], "WORD_COMPARE_LIMIT_EXCEEDED")

    @patch("word.views.MAX_FUZZY_LENGTH", 5, create=True)
    def test_individual_line_length_is_bounded_for_legacy_fuzzy_matching(self):
        response = self.compare({"equal_ratio": .9, "weak_equal_ratio": .7, "output_type": "word"})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data["code"], "WORD_COMPARE_LIMIT_EXCEEDED")

    @patch("word.views.align_paragraphs_indexed", side_effect=RuntimeError("Expensive matching reached"))
    def test_weighted_character_work_is_bounded_before_matching(self, _alignment):
        response = self.compare({"equal_ratio": .9, "weak_equal_ratio": .7, "output_type": "word"},
                                first=word_upload(["ab" * 1000] * 4), second=word_upload(["ba" * 1000] * 4))
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data["code"], "WORD_COMPARE_LIMIT_EXCEEDED")

    @patch("word.views.read_docx_lines_with_index", side_effect=RuntimeError("private-internal-detail"))
    def test_unexpected_processing_failure_is_safe_500(self, _reader):
        response = self.compare({"equal_ratio": .9, "weak_equal_ratio": .7, "output_type": "word"})
        self.assertEqual(response.status_code, 500)
        self.assertEqual(response.data["code"], "INTERNAL_ERROR")
        self.assertNotIn("private-internal-detail", str(response.data))
