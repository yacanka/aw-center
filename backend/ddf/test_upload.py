from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from docx import Document
from rest_framework.test import APIClient

from ddf.models import DDF


def ddf_document(*, repeated_values=False):
    document = Document()
    table = document.add_table(rows=6, cols=6)
    for index, label, value in (
        (0, "Project", "Example"),
        (1, "Document name", "Document"),
        (3, "Commentor", "Reviewer"),
    ):
        table.cell(index, 0).text = label
        table.cell(index, 1).merge(table.cell(index, 5)).text = value
    for cell, value in zip(table.rows[2].cells, ["Number", "A", "Issue", "A" if repeated_values else "B", "Date", "2026-10-13"]):
        cell.text = value
    for cell, value in zip(table.rows[4].cells, ["No", "Authority", "Comment", "Response", "Status", "Reference"]):
        cell.text = value
    for cell, value in zip(table.rows[5].cells, ["1", "Authority", "Repeated", "Repeated", "", ""]):
        cell.text = value
    return document


class DDFUploadTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("ddf-import-user")
        self.user.user_permissions.add(Permission.objects.get(codename="add_ddf"))
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.client.raise_request_exception = False

    def upload(self, document):
        stream = BytesIO()
        document.save(stream)
        return self.upload_bytes(stream.getvalue())

    def upload_bytes(self, content):
        return self.client.post("/api/tools/ddf/upload/", {
            "file": SimpleUploadedFile("ddf.docx", content, content_type="application/octet-stream"),
        }, format="multipart")

    def test_merged_cells_are_collapsed_without_losing_repeated_independent_values(self):
        response = self.upload(ddf_document(repeated_values=True))

        self.assertEqual(response.status_code, 201)
        record = DDF.objects.get(pk=response.data["id"])
        self.assertEqual((record.project, record.doc_name, record.doc_no, record.doc_issue, record.date, record.commentor),
                         ("Example", "Document", "A", "A", "2026-10-13", "Reviewer"))
        self.assertEqual(record.comments, [["1", "Authority", "Repeated", "Repeated", "", ""]])
        self.assertEqual(record.created_by_id, self.user.pk)

    def test_missing_tables_or_short_metadata_rows_return_safe_400_without_persisting(self):
        documents = [Document()]
        short = Document()
        short.add_table(rows=6, cols=2)
        documents.append(short)
        for document in documents:
            with self.subTest(document=document):
                response = self.upload(document)
                self.assertEqual(response.status_code, 400)
                self.assertEqual(response.data["code"], "DDF_IMPORT_INVALID")
        self.assertEqual(DDF.objects.count(), 0)

    def test_empty_metadata_labels_are_rejected_before_persisting(self):
        document = ddf_document()
        document.tables[0].cell(2, 2).text = ""
        response = self.upload(document)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data["code"], "DDF_IMPORT_INVALID")
        self.assertFalse(DDF.objects.exists())

    def test_invalid_office_xml_returns_safe_400(self):
        document = ddf_document()
        original = BytesIO()
        document.save(original)
        rebuilt = BytesIO()
        with ZipFile(original) as source, ZipFile(rebuilt, "w", ZIP_DEFLATED) as destination:
            for entry in source.infolist():
                destination.writestr(entry, b"<broken" if entry.filename == "word/document.xml" else source.read(entry.filename))
        response = self.upload_bytes(rebuilt.getvalue())
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data["code"], "DDF_IMPORT_INVALID")
        self.assertNotIn("broken", str(response.data))
        self.assertFalse(DDF.objects.exists())
